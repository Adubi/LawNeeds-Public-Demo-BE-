import io
import logging
import os
import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional

import fitz
from striprtf.striprtf import rtf_to_text

logger = logging.getLogger(__name__)


@dataclass
class ParsedTable:
    rows: List[List[str]]
    index: int


@dataclass
class ParsedSection:
    heading: Optional[str]
    text: str = ""
    tables: List[ParsedTable] = field(default_factory=list)


@dataclass
class ParsedPage:
    page_number: int
    sections: List[ParsedSection] = field(default_factory=list)


@dataclass
class ParsedLegalDocument:
    document_title: Optional[str]
    pages: List[ParsedPage]
    parser_version: str
    source_name: Optional[str] = None


class StructuredLegalDocumentParser:
    """Parse legal documents into title, pages, sections, headings, and tables."""

    parser_version = "structured-v2"

    def parse(self, file_obj, file_type: str, source_name: Optional[str] = None) -> Dict[str, Any]:
        extension = (file_type or "").lower().lstrip(".")
        if extension == "pdf":
            parsed = self._parse_pdf(file_obj, source_name)
        elif extension == "rtf":
            parsed = self._parse_rtf(file_obj, source_name)
        elif extension == "docx":
            parsed = self._parse_docx(file_obj, source_name)
        else:
            raise ValueError(f"Unsupported structured parser file type: {file_type}")

        return self.to_dict(parsed)

    def to_plain_text(self, parsed_document: Dict[str, Any]) -> str:
        parts: List[str] = []
        title = parsed_document.get("document_title")
        if title:
            parts.append(f"Document: {title}")

        for page in parsed_document.get("pages", []):
            page_number = page.get("page_number")
            parts.append(f"\nPage: {page_number}")
            for section in page.get("sections", []):
                heading = section.get("heading")
                if heading:
                    parts.append(f"\nSection: {heading}")
                text = (section.get("text") or "").strip()
                if text:
                    parts.append(text)
                for table in section.get("tables", []):
                    rows = table.get("rows") or []
                    if rows:
                        parts.append("Table:")
                        parts.extend(self._format_table_rows(rows))

        return "\n".join(part for part in parts if part is not None).strip()

    @staticmethod
    def to_dict(parsed: ParsedLegalDocument) -> Dict[str, Any]:
        return {
            "document_title": parsed.document_title,
            "source_name": parsed.source_name,
            "parser_version": parsed.parser_version,
            "pages": [
                {
                    "page_number": page.page_number,
                    "sections": [
                        {
                            "heading": section.heading,
                            "text": section.text.strip(),
                            "tables": [
                                {
                                    "index": table.index,
                                    "rows": table.rows,
                                }
                                for table in section.tables
                            ],
                        }
                        for section in page.sections
                    ],
                }
                for page in parsed.pages
            ],
        }

    def _parse_pdf(self, file_obj, source_name: Optional[str]) -> ParsedLegalDocument:
        raw = self._read_bytes(file_obj)
        doc = fitz.open(stream=raw, filetype="pdf")
        pages: List[ParsedPage] = []
        first_heading: Optional[str] = None

        for page_index in range(len(doc)):
            page = doc.load_page(page_index)
            lines = self._extract_pdf_lines(page)
            page_tables = self._extract_pdf_tables(page)
            parsed_page = self._lines_to_page(page_index + 1, lines, page_tables)
            pages.append(parsed_page)

            if not first_heading:
                first_heading = self._first_heading(parsed_page)

        return ParsedLegalDocument(
            document_title=(
                self._title_from_pdf_metadata(doc)
                or first_heading
                or self._title_from_source_name(source_name)
            ),
            pages=pages,
            parser_version=self.parser_version,
            source_name=source_name,
        )

    def _extract_pdf_lines(self, page) -> List[Dict[str, Any]]:
        page_dict = page.get_text("dict")
        spans = []
        for block in page_dict.get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                line_text_parts = []
                sizes = []
                fonts = []
                y_values = []
                for span in line.get("spans", []):
                    text = (span.get("text") or "").strip()
                    if not text:
                        continue
                    line_text_parts.append(text)
                    sizes.append(float(span.get("size") or 0))
                    fonts.append(span.get("font") or "")
                    bbox = span.get("bbox") or [0, 0, 0, 0]
                    y_values.append(float(bbox[1]))

                text = " ".join(line_text_parts).strip()
                if text:
                    spans.append(
                        {
                            "text": text,
                            "font_size": max(sizes) if sizes else 0,
                            "font": " ".join(fonts),
                            "y": min(y_values) if y_values else 0,
                        }
                    )

        if not spans:
            return []

        font_sizes = sorted(line["font_size"] for line in spans if line["font_size"])
        body_size = min(font_sizes) if font_sizes else 0
        for line in spans:
            line["is_heading"] = self._looks_like_heading(
                line["text"],
                font_size=line["font_size"],
                body_size=body_size,
                font=line.get("font"),
            )

        return spans

    def _extract_pdf_tables(self, page) -> List[ParsedTable]:
        tables: List[ParsedTable] = []
        if not hasattr(page, "find_tables"):
            return tables

        try:
            found = page.find_tables()
            for index, table in enumerate(found.tables if found else []):
                rows = table.extract()
                clean_rows = self._clean_table_rows(rows or [])
                if clean_rows:
                    tables.append(ParsedTable(rows=clean_rows, index=index))
        except Exception as e:
            logger.debug("PDF table extraction failed: %s", e)

        return tables

    def _parse_rtf(self, file_obj, source_name: Optional[str]) -> ParsedLegalDocument:
        raw = self._read_bytes(file_obj)
        text = rtf_to_text(raw.decode("utf-8", errors="ignore"))
        page = self._plain_text_to_page(text)
        return ParsedLegalDocument(
            document_title=self._first_heading(page) or self._title_from_source_name(source_name),
            pages=[page],
            parser_version=self.parser_version,
            source_name=source_name,
        )

    def _parse_docx(self, file_obj, source_name: Optional[str]) -> ParsedLegalDocument:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise RuntimeError("python-docx is required for DOCX structured parsing") from exc

        self._seek_start(file_obj)
        document = DocxDocument(file_obj)
        sections: List[ParsedSection] = []
        current = ParsedSection(heading=None)
        table_index = 0
        metadata_title = self._clean_title(getattr(document.core_properties, "title", None))

        for paragraph in document.paragraphs:
            text = (paragraph.text or "").strip()
            if not text:
                continue

            style_name = (paragraph.style.name if paragraph.style else "") or ""
            if style_name.lower().startswith("heading") or self._looks_like_heading(text):
                if current.heading or current.text or current.tables:
                    sections.append(current)
                current = ParsedSection(heading=text)
            else:
                current.text = self._append_text(current.text, text)

        for table in document.tables:
            rows = [
                [cell.text.strip() for cell in row.cells]
                for row in table.rows
            ]
            clean_rows = self._clean_table_rows(rows)
            if clean_rows:
                current.tables.append(ParsedTable(rows=clean_rows, index=table_index))
                table_index += 1

        if current.heading or current.text or current.tables:
            sections.append(current)

        page = ParsedPage(page_number=1, sections=sections or [ParsedSection(heading=None, text="")])
        return ParsedLegalDocument(
            document_title=metadata_title or self._first_heading(page) or self._title_from_source_name(source_name),
            pages=[page],
            parser_version=self.parser_version,
            source_name=source_name,
        )

    def _lines_to_page(
        self,
        page_number: int,
        lines: List[Dict[str, Any]],
        tables: List[ParsedTable],
    ) -> ParsedPage:
        sections: List[ParsedSection] = []
        current = ParsedSection(heading=None)

        for line in lines:
            text = line["text"].strip()
            if not text:
                continue

            if line.get("is_heading"):
                if current.heading or current.text or current.tables:
                    sections.append(current)
                current = ParsedSection(heading=text)
            else:
                current.text = self._append_text(current.text, text)

        if tables:
            current.tables.extend(tables)

        if current.heading or current.text or current.tables:
            sections.append(current)

        return ParsedPage(page_number=page_number, sections=sections or [ParsedSection(heading=None, text="")])

    def _plain_text_to_page(self, text: str) -> ParsedPage:
        sections: List[ParsedSection] = []
        current = ParsedSection(heading=None)

        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                continue

            if self._looks_like_heading(line):
                if current.heading or current.text:
                    sections.append(current)
                current = ParsedSection(heading=line)
            else:
                current.text = self._append_text(current.text, line)

        if current.heading or current.text:
            sections.append(current)

        return ParsedPage(page_number=1, sections=sections or [ParsedSection(heading=None, text=text.strip())])

    @staticmethod
    def _append_text(existing: str, addition: str) -> str:
        if not existing:
            return addition
        return f"{existing}\n{addition}"

    @staticmethod
    def _read_bytes(file_obj) -> bytes:
        if isinstance(file_obj, bytes):
            return file_obj
        if isinstance(file_obj, bytearray):
            return bytes(file_obj)
        StructuredLegalDocumentParser._seek_start(file_obj)
        raw = file_obj.read()
        StructuredLegalDocumentParser._seek_start(file_obj)
        return raw

    @staticmethod
    def _seek_start(file_obj) -> None:
        if hasattr(file_obj, "seek"):
            file_obj.seek(0)

    @staticmethod
    def _looks_like_heading(
        text: str,
        font_size: Optional[float] = None,
        body_size: Optional[float] = None,
        font: Optional[str] = None,
    ) -> bool:
        clean = text.strip()
        if not clean or len(clean) > 180:
            return False
        if re.fullmatch(r"\d+", clean):
            return False
        if font_size and body_size and font_size >= body_size * 1.18:
            return True
        if font and "bold" in font.lower() and len(clean.split()) <= 14:
            return True
        if re.match(r"^(\d+(\.\d+)*|[A-Z])[\).]\s+\S+", clean) and len(clean.split()) <= 18:
            return True
        letters = [char for char in clean if char.isalpha()]
        if letters and sum(1 for char in letters if char.isupper()) / len(letters) > 0.75 and len(clean.split()) <= 14:
            return True
        if clean.endswith(":") and len(clean.split()) <= 14:
            return True
        return False

    @staticmethod
    def _first_heading(page: ParsedPage) -> Optional[str]:
        for section in page.sections:
            heading = (section.heading or "").strip()
            if heading:
                return heading
        return None

    @staticmethod
    def _title_from_pdf_metadata(doc) -> Optional[str]:
        metadata = getattr(doc, "metadata", {}) or {}
        return StructuredLegalDocumentParser._clean_title(metadata.get("title"))

    @staticmethod
    def _title_from_source_name(source_name: Optional[str]) -> Optional[str]:
        if not source_name:
            return None
        file_name = os.path.basename(source_name)
        title = os.path.splitext(file_name)[0].replace("_", " ").strip()
        return StructuredLegalDocumentParser._clean_title(title)

    @staticmethod
    def _clean_title(value: Optional[str]) -> Optional[str]:
        title = (value or "").strip()
        if not title:
            return None
        lower = title.lower()
        if lower in {"untitled", "document", "microsoft word - document"}:
            return None
        if len(title) <= 2:
            return None
        return title

    @staticmethod
    def _clean_table_rows(rows: List[List[Any]]) -> List[List[str]]:
        cleaned: List[List[str]] = []
        for row in rows:
            cleaned_row = [str(cell or "").strip() for cell in row]
            if any(cleaned_row):
                cleaned.append(cleaned_row)
        return cleaned

    @staticmethod
    def _format_table_rows(rows: List[List[str]]) -> List[str]:
        return [" | ".join(str(cell or "").strip() for cell in row) for row in rows]
